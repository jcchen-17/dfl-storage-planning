"""Build the historical IEEE13/Smart-DS dataset consumed by storage_dfl.

The raw directory is read-only by convention.  All derived artifacts are written
under data/processed/ieee13_smartds_dfl.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import tarfile
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
import pyarrow as pa
import pyarrow.csv as pacsv
import pyarrow.parquet as pq
from openpyxl import load_workbook


BUSES = ("650", "632", "633", "634", "645", "646", "671", "680", "684", "611", "652", "692", "675")
PHASES = ("A", "B", "C")
BUS_INDEX = {bus: i for i, bus in enumerate(BUSES)}
PHASE_INDEX = {phase: i for i, phase in enumerate(PHASES)}
YEARS = (2016, 2017, 2018)
SPLITS = {2016: "train", 2017: "validation", 2018: "test"}

# Official IEEE13 branches after collapsing artificial bus 670 into branch 632--671.
BRANCHES = (
    ("reg_650_632", "650", "632", "regulator", "ABC", 2000, "mtx601", 4.16, 4.16),
    ("line_632_633", "632", "633", "line", "ABC", 500, "mtx602", 4.16, 4.16),
    ("xfm_633_634", "633", "634", "transformer", "ABC", 0, "XFM1", 4.16, 0.48),
    ("line_632_645", "632", "645", "line", "BC", 500, "mtx603", 4.16, 4.16),
    ("line_645_646", "645", "646", "line", "BC", 300, "mtx603", 4.16, 4.16),
    ("line_632_671", "632", "671", "line_collapsed_670", "ABC", 2000, "mtx601", 4.16, 4.16),
    ("line_671_680", "671", "680", "line", "ABC", 1000, "mtx601", 4.16, 4.16),
    ("line_671_684", "671", "684", "line", "AC", 300, "mtx604", 4.16, 4.16),
    ("line_684_611", "684", "611", "line", "C", 300, "mtx605", 4.16, 4.16),
    ("line_684_652", "684", "652", "line", "A", 800, "mtx607", 4.16, 4.16),
    ("switch_671_692", "671", "692", "switch_closed", "ABC", 0, "switch", 4.16, 4.16),
    ("line_692_675", "692", "675", "line", "ABC", 500, "mtx606", 4.16, 4.16),
)

# element, physical bus, model bus, connection, terminal phases, kW, kvar, model.
LOADS = (
    ("671", "671", "671", "delta", "ABC", 1155.0, 660.0, 1),
    ("634a", "634", "634", "wye", "A", 160.0, 110.0, 1),
    ("634b", "634", "634", "wye", "B", 120.0, 90.0, 1),
    ("634c", "634", "634", "wye", "C", 120.0, 90.0, 1),
    ("645", "645", "645", "wye", "B", 170.0, 125.0, 1),
    ("646", "646", "646", "delta", "BC", 230.0, 132.0, 2),
    ("692", "692", "692", "delta", "CA", 170.0, 151.0, 5),
    ("675a", "675", "675", "wye", "A", 485.0, 190.0, 1),
    ("675b", "675", "675", "wye", "B", 68.0, 60.0, 1),
    ("675c", "675", "675", "wye", "C", 290.0, 212.0, 1),
    ("611", "611", "611", "wye", "C", 170.0, 80.0, 5),
    ("652", "652", "652", "wye", "A", 128.0, 86.0, 2),
    ("670a", "670", "671", "wye", "A", 17.0, 10.0, 1),
    ("670b", "670", "671", "wye", "B", 66.0, 38.0, 1),
    ("670c", "670", "671", "wye", "C", 117.0, 68.0, 1),
)

PV = (("pv_634", "634", "ABC", 0.25, 15), ("pv_675", "675", "ABC", 0.65, 25), ("pv_680", "680", "ABC", 0.75, 25))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workspace", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--azure-sample-functions", type=int, default=5000)
    parser.add_argument("--skip-wide-csv", action="store_true")
    return parser.parse_args()


def write_csv(path: Path, header: list[str], rows) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def node_rows():
    phase_map = {"650": "ABC", "632": "ABC", "633": "ABC", "634": "ABC", "645": "BC", "646": "BC", "671": "ABC", "680": "ABC", "684": "AC", "611": "C", "652": "A", "692": "ABC", "675": "ABC"}
    voltage = {bus: (0.48 if bus == "634" else 4.16) for bus in BUSES}
    base_p = {bus: 0.0 for bus in BUSES}
    base_q = {bus: 0.0 for bus in BUSES}
    for _, _, model_bus, _, _, kw, kvar, _ in LOADS:
        base_p[model_bus] += kw / 1000.0
        base_q[model_bus] += kvar / 1000.0
    pv_cap = {bus: 0.0 for bus in BUSES}
    for _, bus, _, cap, _ in PV:
        pv_cap[bus] += cap
    for bus in BUSES:
        yield [bus, phase_map[bus], voltage[bus], base_p[bus], base_q[bus], pv_cap[bus], int(bus in {"632", "671", "675", "680"}), int(bus == "675"), int(bus == "671")]


def load_components():
    """Split connection totals into the nodal phase approximation used by DFL."""
    components = []
    for name, physical_bus, model_bus, conn, terminals, kw, kvar, model in LOADS:
        phase_list = list(terminals)
        share = 1.0 / len(phase_list)
        for phase in phase_list:
            components.append(
                {
                    "component": f"{name}_{phase}", "element": name, "physical_bus": physical_bus,
                    "model_bus": model_bus, "phase": phase, "connection": conn,
                    "terminals": terminals, "base_mw": kw * share / 1000.0,
                    "base_mvar": kvar * share / 1000.0, "load_model": model,
                }
            )
    return components


def read_smartds_loads(raw: Path, year: int):
    folder = raw / "smartds" / str(year) / "GSO" / "industrial" / "load_data"
    profiles = []
    for path in sorted(folder.glob("*.parquet")):
        table = pq.read_table(path, columns=["total_site_electricity_kw", "total_site_electricity_kvar"])
        p = np.asarray(table.column(0).combine_chunks().to_numpy(), dtype=np.float64)
        q = np.asarray(table.column(1).combine_chunks().to_numpy(), dtype=np.float64)
        if len(p) != 35040:
            raise ValueError(f"{path}: expected 35040 quarter-hours, got {len(p)}")
        p = np.nanmean(p.reshape(8760, 4), axis=1)
        q = np.nanmean(q.reshape(8760, 4), axis=1)
        p_scale = max(float(np.nanpercentile(np.maximum(p, 0.0), 99.5)), 1e-6)
        q_scale = max(float(np.nanpercentile(np.maximum(q, 0.0), 99.5)), 1e-6)
        load_factor = float(np.nanmean(np.maximum(p, 0.0)) / p_scale)
        profiles.append({"id": path.stem, "path": path, "p": np.maximum(p, 0.0) / p_scale, "q": np.maximum(q, 0.0) / q_scale, "load_factor": load_factor})
    return sorted(profiles, key=lambda item: (item["load_factor"], item["id"]))


def map_loads(raw: Path, year: int, components: list[dict]):
    profiles = read_smartds_loads(raw, year)
    p = np.zeros((8760, len(BUSES), 3), dtype=np.float64)
    q = np.zeros_like(p)
    mapping = []
    for index, component in enumerate(components):
        # Spread components evenly over the load-factor-ordered profile library.
        profile_index = int(round(index * (len(profiles) - 1) / max(1, len(components) - 1)))
        profile = profiles[profile_index]
        bi, pi = BUS_INDEX[component["model_bus"]], PHASE_INDEX[component["phase"]]
        p[:, bi, pi] += component["base_mw"] * profile["p"]
        q[:, bi, pi] += component["base_mvar"] * profile["q"]
        mapping.append([
            year, component["component"], component["element"], component["physical_bus"],
            component["model_bus"], component["phase"], component["connection"], component["terminals"],
            component["base_mw"], component["base_mvar"], profile["id"], profile["load_factor"],
            "P/Q各按Smart-DS曲线99.5分位归一化；三角形负荷按端子相均分为节点相等效注入",
        ])
    # Treat official totals as the planning feeder peak ratings.
    p_factor = 3.466 / float(np.max(np.sum(p, axis=(1, 2))))
    q_factor = 2.102 / float(np.max(np.sum(q, axis=(1, 2))))
    p *= p_factor
    q *= q_factor
    return p, q, mapping, p_factor, q_factor


def read_solar_weather(raw: Path, year: int):
    base = raw / "smartds" / str(year) / "GSO" / "industrial" / "solar_data"
    result = {}
    for tilt in (15, 25):
        path = base / f"GSO_36.0976_-80.0003_{tilt}_180_full.csv"
        table = pacsv.read_csv(path)
        pv = np.asarray(table["kW Generated (1000 kW Array)"].to_numpy(), dtype=float)
        temp = np.asarray(table["Temperature"].to_numpy(), dtype=float)
        if len(pv) != 35040:
            raise ValueError(f"{path}: expected 35040 rows, got {len(pv)}")
        result[tilt] = np.maximum(pv.reshape(8760, 4).mean(axis=1) / 1000.0, 0.0)
        if tilt == 15:
            result["temperature_c"] = temp.reshape(8760, 4).mean(axis=1)
    return result


def make_pv(weather: dict):
    array = np.zeros((8760, len(BUSES), 3), dtype=np.float64)
    for _, bus, phases, capacity, tilt in PV:
        for phase in phases:
            array[:, BUS_INDEX[bus], PHASE_INDEX[phase]] = capacity * weather[tilt] / len(phases)
    return array


def read_azure_workload(raw: Path, sample_functions: int):
    archive = raw / "data_center" / "azure_functions_2019" / "azurefunctions_dataset2019.tar.xz"
    hourly, sampled = [], []
    with tarfile.open(archive, "r:xz") as tar:
        for day in range(1, 15):
            member = tar.getmember(f"invocations_per_function_md.anon.d{day:02d}.csv")
            totals = np.zeros(1440, dtype=np.float64)
            rows = 0
            with tar.extractfile(member) as stream:
                next(stream)
                for line in stream:
                    if rows >= sample_functions:
                        break
                    start = 0
                    for _ in range(4):
                        start = line.find(b",", start) + 1
                    values = np.fromstring(line[start:].decode("ascii"), sep=",")
                    if values.size == 1440:
                        totals += values
                        rows += 1
            hourly.extend(totals.reshape(24, 60).sum(axis=1))
            sampled.append(rows)
    values = np.log1p(np.asarray(hourly, dtype=np.float64))
    low, high = np.percentile(values, [5, 95])
    normalized = 0.20 + 0.60 * np.clip((values - low) / max(high - low, 1e-9), 0.0, 1.0)
    return normalized, sampled


def expand_workload(base_14d: np.ndarray, year: int):
    # Preserve the measured 14-day intraday ordering; rotate the starting day by year.
    shift = (year - 2016) * 24 * 3
    indices = (np.arange(8760) + shift) % len(base_14d)
    return base_14d[indices]


def read_duke_carbon(raw: Path):
    path = raw / "grid" / "carbon" / "eia930" / "DUK.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Published Hourly Data"]
    result = {2021: [], 2022: [], 2023: []}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        timestamp, intensity = row[1], row[88]
        if isinstance(timestamp, datetime) and timestamp.year in result and intensity is not None:
            result[timestamp.year].append((timestamp, float(intensity) * 0.45359237))  # lb/kWh -> t/MWh
    workbook.close()
    arrays = {}
    for year, rows in result.items():
        rows.sort(key=lambda item: item[0])
        if len(rows) != 8760:
            raise ValueError(f"EIA DUK {year}: expected 8760 non-missing hours, got {len(rows)}")
        arrays[year] = np.asarray([value for _, value in rows], dtype=np.float64)
    return arrays


def tariff_series(raw: Path, timestamps: list[datetime]):
    path = raw / "grid" / "price" / "openei_usurdb" / "OPT-V_2026.json"
    item = json.loads(path.read_text(encoding="utf-8"))["items"][0]
    prices = np.empty(len(timestamps), dtype=np.float64)
    periods = np.empty(len(timestamps), dtype=np.int16)
    for i, timestamp in enumerate(timestamps):
        schedule = item["energyweekdayschedule"] if timestamp.weekday() < 5 else item["energyweekendschedule"]
        period = int(schedule[timestamp.month - 1][timestamp.hour])
        component = item["energyratestructure"][period][0]
        prices[i] = 1000.0 * (float(component.get("rate", 0.0)) + float(component.get("adj", 0.0)))
        periods[i] = period
    return prices, periods, item


def arrow_table(columns: dict[str, np.ndarray | list]):
    return pa.table({name: pa.array(values) for name, values in columns.items()})


def main() -> None:
    args = parse_args()
    workspace = args.workspace.resolve()
    raw = workspace / "data" / "raw"
    output = workspace / "data" / "processed" / "ieee13_smartds_dfl"
    output.mkdir(parents=True, exist_ok=True)

    write_csv(output / "ieee13_nodes.csv", ["bus", "phases", "kv_ll", "base_load_mw", "base_load_mvar", "pv_capacity_mw", "bess_candidate", "data_center", "generator"], node_rows())
    write_csv(output / "ieee13_branches.csv", ["branch", "from_bus", "to_bus", "type", "phases", "length_ft", "linecode", "from_kv_ll", "to_kv_ll"], BRANCHES)
    components = load_components()
    write_csv(output / "ieee13_load_connections.csv", ["component", "element", "physical_bus", "model_bus", "phase", "connection", "terminals", "base_mw", "base_mvar", "load_model"], ([c[k] for k in ("component", "element", "physical_bus", "model_bus", "phase", "connection", "terminals", "base_mw", "base_mvar", "load_model")] for c in components))

    azure, sampled = read_azure_workload(raw, args.azure_sample_functions)
    carbon = read_duke_carbon(raw)
    carbon_year = {2016: 2021, 2017: 2022, 2018: 2023}
    all_columns: dict[str, list | np.ndarray] = {}
    accumulated: dict[str, list[np.ndarray]] = {}
    mapping_rows, scale_rows = [], []

    for year in YEARS:
        timestamps = [datetime(year, 1, 1) + timedelta(hours=i) for i in range(8760)]
        load_p, load_q, rows, p_factor, q_factor = map_loads(raw, year, components)
        mapping_rows.extend(rows)
        scale_rows.append([year, p_factor, q_factor, float(load_p.sum(axis=(1, 2)).max()), float(load_q.sum(axis=(1, 2)).max())])
        weather = read_solar_weather(raw, year)
        pv = make_pv(weather)
        workload = expand_workload(azure, year)
        pue = np.clip(1.15 + 0.006 * np.maximum(weather["temperature_c"] - 18.0, 0.0), 1.15, 1.35)
        price, price_period, tariff = tariff_series(raw, timestamps)
        globals_for_year = {
            "timestamp": np.asarray(timestamps, dtype="datetime64[ms]"),
            "year": np.full(8760, year, dtype=np.int16),
            "split": np.asarray([SPLITS[year]] * 8760),
            "month": np.asarray([t.month for t in timestamps], dtype=np.int8),
            "day": np.asarray([t.day for t in timestamps], dtype=np.int8),
            "hour": np.asarray([t.hour for t in timestamps], dtype=np.int8),
            "day_of_year": np.asarray([t.timetuple().tm_yday for t in timestamps], dtype=np.int16),
            "day_of_week": np.asarray([t.weekday() for t in timestamps], dtype=np.int8),
            "is_weekend": np.asarray([t.weekday() >= 5 for t in timestamps], dtype=np.int8),
            "temperature_c": weather["temperature_c"],
            "workload_arrival": workload,
            "pue": pue,
            "grid_price_per_mwh": price,
            "tariff_energy_period": price_period,
            "grid_carbon_t_per_mwh": carbon[carbon_year[year]],
            "grid_available": np.ones(8760, dtype=np.int8),
            "total_load_mw": load_p.sum(axis=(1, 2)),
            "total_load_mvar": load_q.sum(axis=(1, 2)),
            "total_pv_available_mw": pv.sum(axis=(1, 2)),
        }
        year_columns = dict(globals_for_year)
        for bi, bus in enumerate(BUSES):
            year_columns[f"active_load_mw__{bus}"] = load_p[:, bi, :].sum(axis=1)
            year_columns[f"reactive_load_mvar__{bus}"] = load_q[:, bi, :].sum(axis=1)
            year_columns[f"pv_available_mw__{bus}"] = pv[:, bi, :].sum(axis=1)
            for pi, phase in enumerate(PHASES):
                year_columns[f"active_load_mw__{bus}__{phase}"] = load_p[:, bi, pi]
                year_columns[f"reactive_load_mvar__{bus}__{phase}"] = load_q[:, bi, pi]
                year_columns[f"pv_available_mw__{bus}__{phase}"] = pv[:, bi, pi]
        for name, values in year_columns.items():
            accumulated.setdefault(name, []).append(np.asarray(values))

    all_columns = {name: np.concatenate(parts) for name, parts in accumulated.items()}
    table = arrow_table(all_columns)
    pq.write_table(table, output / "dfl_hourly_wide.parquet", compression="zstd")
    if not args.skip_wide_csv:
        pacsv.write_csv(table, output / "dfl_hourly_wide.csv")

    global_dictionary = {
        "timestamp": ("模型小时起点", "datetime", "重建的Smart-DS标准年小时轴", "索引，不进入轨迹"),
        "year": ("数据年份", "year", "Smart-DS", "用于划分数据"),
        "split": ("训练/验证/测试标签", "category", "2016=train, 2017=validation, 2018=test", "用于划分数据"),
        "month": ("月份1–12", "month", "timestamp派生", "可作为条件"),
        "day": ("月内日期", "day", "timestamp派生", "可作为条件"),
        "hour": ("小时0–23", "hour", "timestamp派生", "可作为条件"),
        "day_of_year": ("年内日序1–365", "day", "timestamp派生", "生成季节sin/cos条件"),
        "day_of_week": ("星期序号：0周一，6周日", "category", "timestamp派生", "生成周末条件"),
        "is_weekend": ("周末标记：周六/周日为1", "0/1", "timestamp派生", "生成weekend_fraction条件"),
        "temperature_c": ("室外气温", "degC", "Smart-DS GSO solar/weather", "生成PUE及CVAE条件"),
        "workload_arrival": ("归一化数据中心计算任务到达量，不是MW", "p.u.", "Azure Functions调用量", "Scenario/CVAE轨迹"),
        "pue": ("数据中心总功率/IT功率", "ratio", "由temperature_c按假设公式派生", "Scenario/CVAE轨迹"),
        "grid_price_per_mwh": ("上级电网边际电量价格，不含需量费", "USD/MWh", "OpenEI OPT-V 2026", "Scenario/CVAE及运行成本"),
        "tariff_energy_period": ("OPT-V内部时段编号0/1/2", "category", "OpenEI费率日历", "解释电价，当前不直接入codec"),
        "grid_carbon_t_per_mwh": ("消费电力平均碳排放强度", "tCO2/MWh", "EIA-930 DUK", "Scenario/CVAE及碳约束"),
        "grid_available": ("上级电网可用率；历史正常场景均为1", "0/1", "场景假设", "Scenario；极端场景可改为0"),
        "total_load_mw": ("13节点普通工业负荷有功合计，不含数据中心", "MW", "节点有功求和", "校验/分析"),
        "total_load_mvar": ("13节点普通工业负荷无功合计，不含数据中心", "Mvar", "节点无功求和", "校验/分析"),
        "total_pv_available_mw": ("三个PV节点最大可用有功合计，削减前", "MW", "节点PV可用量求和", "校验/分析"),
    }
    dictionary_rows = []
    for position, name in enumerate(table.column_names, start=1):
        if name in global_dictionary:
            meaning, unit, source, usage = global_dictionary[name]
            bus = phase = ""
        else:
            parts = name.split("__")
            variable, bus = parts[0], parts[1]
            phase = parts[2] if len(parts) == 3 else "ALL"
            labels = {
                "active_load_mw": ("普通工业负荷有功", "MW", "Smart-DS负荷曲线映射并按IEEE13容量缩放", "Scenario/CVAE输入"),
                "reactive_load_mvar": ("普通工业负荷无功", "Mvar", "Smart-DS无功曲线映射并按IEEE13容量缩放", "Scenario规划输入；当前codec未生成"),
                "pv_available_mw": ("削减前PV最大可用有功", "MW", "Smart-DS太阳能曲线按安装容量缩放", "Scenario/CVAE输入"),
            }
            base_meaning, unit, source, usage = labels[variable]
            meaning = f"节点{bus}{'总' if phase == 'ALL' else phase + '相'}{base_meaning}"
        dictionary_rows.append([position, name, meaning, unit, source, bus, phase, usage])
    write_csv(
        output / "dfl_hourly_wide_dictionary.csv",
        ["column_position", "column_name", "meaning", "unit", "source_or_derivation", "bus", "phase", "dfl_usage"],
        dictionary_rows,
    )

    write_csv(output / "smartds_to_ieee13_mapping.csv", ["year", "component", "ieee_element", "physical_bus", "model_bus", "phase", "connection", "terminals", "base_mw", "base_mvar", "smartds_profile", "smartds_load_factor", "method"], mapping_rows)
    write_csv(output / "annual_scaling.csv", ["year", "p_scale_factor", "q_scale_factor", "result_peak_mw", "result_peak_mvar"], scale_rows)

    # Create Scenario-compatible 120-hour arrays.  One window starts every 24 h.
    bus_p = np.stack([all_columns[f"active_load_mw__{bus}"] for bus in BUSES], axis=1)
    bus_q = np.stack([all_columns[f"reactive_load_mvar__{bus}"] for bus in BUSES], axis=1)
    bus_pv = np.stack([all_columns[f"pv_available_mw__{bus}"] for bus in BUSES], axis=1)
    phase_p = np.stack([all_columns[f"active_load_mw__{bus}__{phase}"] for bus in BUSES for phase in PHASES], axis=1).reshape(-1, len(BUSES), 3)
    phase_q = np.stack([all_columns[f"reactive_load_mvar__{bus}__{phase}"] for bus in BUSES for phase in PHASES], axis=1).reshape(-1, len(BUSES), 3)
    phase_pv = np.stack([all_columns[f"pv_available_mw__{bus}__{phase}"] for bus in BUSES for phase in PHASES], axis=1).reshape(-1, len(BUSES), 3)
    starts, index_rows = [], []
    for yi, year in enumerate(YEARS):
        offset = yi * 8760
        for local_start in range(0, 8760 - 120 + 1, 24):
            starts.append(offset + local_start)
            ts = datetime(year, 1, 1) + timedelta(hours=local_start)
            index_rows.append([len(starts) - 1, f"{year}_{local_start // 24:03d}", year, SPLITS[year], ts.isoformat(), local_start, local_start + 120])
    starts = np.asarray(starts, dtype=np.int32)
    def windows(values):
        return np.stack([values[start:start + 120] for start in starts]).astype(np.float32)
    context = np.empty((len(starts), 4), dtype=np.float32)
    for i, start in enumerate(starts):
        ts = all_columns["timestamp"][start].astype("datetime64[ms]").astype(datetime)
        angle = 2.0 * math.pi * (ts.timetuple().tm_yday - 1) / 365.0
        context[i] = [math.sin(angle), math.cos(angle), float(np.mean(all_columns["is_weekend"][start:start + 120])), float(np.mean(all_columns["temperature_c"][start:start + 120]))]
    np.savez_compressed(
        output / "dfl_training_windows_120h.npz",
        scenario_name=np.asarray([row[1] for row in index_rows]), split=np.asarray([row[3] for row in index_rows]),
        context=context, active_load_mw=windows(bus_p), reactive_load_mvar=windows(bus_q),
        pv_available_mw=windows(bus_pv), active_load_phase_mw=windows(phase_p),
        reactive_load_phase_mvar=windows(phase_q), pv_available_phase_mw=windows(phase_pv),
        workload_arrival=windows(all_columns["workload_arrival"]), pue=windows(all_columns["pue"]),
        grid_price_per_mwh=windows(all_columns["grid_price_per_mwh"]),
        grid_carbon_t_per_mwh=windows(all_columns["grid_carbon_t_per_mwh"]),
        grid_available=windows(all_columns["grid_available"]), buses=np.asarray(BUSES), phases=np.asarray(PHASES),
        context_names=np.asarray(["start_doy_sin", "start_doy_cos", "weekend_fraction", "mean_temperature_c"]),
    )
    write_csv(output / "dfl_window_index.csv", ["window_id", "scenario_name", "source_year", "split", "start_timestamp", "start_hour", "end_hour_exclusive"], index_rows)

    compatibility = [
        ["Smart-DS GSO industrial", "load/PV/weather", "YES-scaled", "3年15分钟原始曲线聚合为小时；P/Q按IEEE13额定峰值缩放；PV按0.25/0.65/0.75 MW容量缩放", "主要历史场景源"],
        ["Azure Functions 2019", "data-center workload", "YES-proxy", f"14天分钟调用数；每天确定性抽样前{args.azure_sample_functions}个匿名函数，汇总到小时并循环扩展", "不是同地点同年份，只能作为工作量形状代理"],
        ["EIA-930 DUK", "grid carbon", "YES-region/time relabel", "取2021/2022/2023连续UTC小时分别映射到2016/2017/2018；lb/kWh乘0.45359237转t/MWh", "DUK与OPT-V同属Duke区域，但与Smart-DS不是同步观测"],
        ["OpenEI USURDB OPT-V 2026", "energy price", "YES-tariff", "按月/工作日/小时费率表生成边际电价；$/kWh乘1000转$/MWh", "当前DFL只用电量价，需量电费另列但未进入逐小时目标"],
        ["IEEE PES IEEE13", "network/topology", "YES-static", "保留节点相别、连接制与支路；670虚拟负荷点折算到671", "网络额定数据，不是历史时间序列"],
    ]
    write_csv(output / "source_compatibility.csv", ["source", "role", "compatibility", "processing", "limitation"], compatibility)

    schema_rows = [
        ["context", "[N,4]", "float32", "CVAE condition: season sin/cos, weekend share, mean temperature"],
        ["active_load_mw", "[N,120,13]", "float32", "current Scenario/codec input"],
        ["reactive_load_mvar", "[N,120,13]", "float32", "current planning input; codec currently does not generate it"],
        ["pv_available_mw", "[N,120,13]", "float32", "current Scenario/codec input"],
        ["active_load_phase_mw", "[N,120,13,3]", "float32", "future three-phase model input"],
        ["reactive_load_phase_mvar", "[N,120,13,3]", "float32", "future three-phase model input"],
        ["pv_available_phase_mw", "[N,120,13,3]", "float32", "future three-phase model input"],
        ["workload_arrival", "[N,120]", "float32", "normalized Azure arrival workload"],
        ["pue", "[N,120]", "float32", "temperature-derived PUE"],
        ["grid_price_per_mwh", "[N,120]", "float32", "OPT-V energy price only"],
        ["grid_carbon_t_per_mwh", "[N,120]", "float32", "EIA-930 DUK consumed-electricity intensity"],
        ["grid_available", "[N,120]", "float32", "normal historical data = 1; outage extremes added later"],
    ]
    write_csv(output / "dfl_dataset_schema.csv", ["field", "shape", "dtype", "meaning"], schema_rows)

    validation = {
        "hours": len(all_columns["year"]), "windows": len(starts), "window_hours": 120,
        "windows_by_split": {split: int(np.sum(np.asarray([row[3] for row in index_rows]) == split)) for split in SPLITS.values()},
        "buses": list(BUSES), "phases": list(PHASES), "official_peak_load_mw": 3.466,
        "official_peak_load_mvar": 2.102, "pv_capacity_mw": 1.65, "data_center_bus": "675",
        "data_center_model_peak_mw_approx": 0.60, "storage_candidates": ["632", "671", "675", "680"],
        "azure_functions_sampled_per_day": sampled,
        "ranges": {name: [float(np.min(all_columns[name])), float(np.max(all_columns[name]))] for name in ("total_load_mw", "total_load_mvar", "total_pv_available_mw", "workload_arrival", "pue", "grid_price_per_mwh", "grid_carbon_t_per_mwh")},
        "sha256_windows": hashlib.sha256((output / "dfl_training_windows_120h.npz").read_bytes()).hexdigest(),
    }
    (output / "validation_summary.json").write_text(json.dumps(validation, indent=2, ensure_ascii=False), encoding="utf-8")

    readme = f"""# IEEE13 × Smart-DS × DFL 历史训练集

本目录含 {len(all_columns['year']):,} 个连续小时和 {len(starts):,} 个 120 小时场景窗口。2016/2017/2018 分别固定为 train/validation/test，每天滑动一次，因此每年 361 个窗口。

## 当前 DFL 真正读取的形式

`Scenario` 需要节点级 `P[120,13]`、`Q[120,13]`、`PV[120,13]`，以及 workload、PUE、电价、碳强度、可用率各 `[120]`，条件向量为 `[4]`。这些已经写入 `dfl_training_windows_120h.npz`。同时保存了 `[120,13,3]` 三相数组，供后续把平衡 LinDistFlow 改为三相不平衡模型；当前 `ScenarioCodec` 尚未生成 Q，也尚未使用 outage，这是代码下一步应改的地方。

## 数据融合边界

这不是一个“同一地点、同一时钟”的联合观测数据集。Smart-DS 提供园区负荷/PV/气温；Azure 只提供 14 天匿名云函数工作量形状；EIA 2021–2023 碳强度按连续小时重标到 2016–2018；OPT-V 是 2026 费率规则。它们在量纲、区域用途和规划含义上可融合，但不能据此声称真实相关性。CVAE训练时应做联合敏感性/错配检验。

## 关键文件

- `dfl_hourly_wide.parquet/csv`：一行一小时的大表，含节点总量和节点—相量。
- `dfl_training_windows_120h.npz`：可直接转成现有 `Scenario` 的训练张量。
- `ieee13_nodes.csv`、`ieee13_branches.csv`、`ieee13_load_connections.csv`：静态拓扑和连接。
- `smartds_to_ieee13_mapping.csv`：逐年逐负荷相的 profile 映射依据。
- `source_compatibility.csv`：四类外部数据能否匹配以及限制。
- `validation_summary.json`：规模、范围、峰值和文件哈希。

原始文件未被修改。PUE 假设为 `clip(1.15 + 0.006*max(T-18,0), 1.15, 1.35)`；正常历史场景 `grid_available=1`，停电/极端事件应在后续极端场景模块中单独注入。
"""
    (output / "README.md").write_text(readme, encoding="utf-8")
    print(json.dumps(validation, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
